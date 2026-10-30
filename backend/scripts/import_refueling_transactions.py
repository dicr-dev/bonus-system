"""Import refuelling transactions from an Excel report into CargoRun.

The script is intentionally local: CargoRun accepts the user's browser/network
session, while a server-side import is not guaranteed to have that access.
"""

from __future__ import annotations

import argparse
import getpass
import hashlib
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

BASE_URL = "https://delko.cargorun.ru"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/155.0.0.0 Safari/537.36"
)
REQUIRED_COLUMNS = {
    "Дата заправки",
    "Поставщик",
    "ТС",
    "Топливная карта",
    "АЗС",
    "Адрес",
    "Количество, л.",
    "Сумма",
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def parse_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, str):
        return datetime.strptime(value.strip(), "%d.%m.%Y %H:%M:%S").isoformat()  # noqa: DTZ007
    raise ValueError("Дата заправки должна быть датой или строкой ДД.ММ.ГГГГ ЧЧ:ММ:СС.")


def number(value: Any, column_name: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"В колонке «{column_name}» должно быть число.") from error


def report_rows(path: Path, limit: int) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator, ())]
    index = {header: position for position, header in enumerate(headers) if header}
    missing = REQUIRED_COLUMNS - set(index)
    if missing:
        raise ValueError(f"Не найдены обязательные колонки: {', '.join(sorted(missing))}.")

    rows: list[dict[str, Any]] = []
    for excel_row, row in enumerate(iterator, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        try:
            item = {
                "excel_row": excel_row,
                "created_at": parse_date(row[index["Дата заправки"]]),
                "device_number": clean(row[index["ТС"]]),
                "card_number": clean(row[index["Топливная карта"]]),
                "supplier": clean(row[index["Поставщик"]]),
                "station": clean(row[index["АЗС"]]),
                "address": clean(row[index["Адрес"]]),
                "volume": number(row[index["Количество, л."]], "Количество, л."),
                "cost": number(row[index["Сумма"]], "Сумма"),
            }
        except (IndexError, TypeError) as error:
            raise ValueError(f"Строка Excel {excel_row}: неполные данные.") from error
        if not item["address"]:
            raise ValueError(f"Строка Excel {excel_row}: не указан адрес АЗС.")
        rows.append(item)
        if len(rows) == limit:
            break
    if not rows:
        raise ValueError("В файле нет строк для импорта.")
    return rows


def client(token: str, device_id: str) -> httpx.Client:
    return httpx.Client(
        base_url=BASE_URL,
        timeout=30,
        headers={
            "Authorization": f"Bearer {token}",
            "Referer": f"{BASE_URL}/",
            "X-Device-Id": device_id,
            "X-Requested-With": "XMLHttpRequest",
            "User-Agent": USER_AGENT,
            "Content-Type": "application/json; charset=utf-8",
        },
    )


def coordinates(api: httpx.Client, address: str) -> tuple[float, float]:
    response = api.post("/api/Map/GetMixedAddresses", json={"address": address, "count": 1})
    response.raise_for_status()
    suggestions = response.json().get("suggestions") or []
    for suggestion in suggestions:
        point = (suggestion.get("location") or {}).get("coordinates") or []
        if len(point) == 2:
            return float(point[0]), float(point[1])
    raise ValueError("Не удалось определить координаты АЗС по адресу из файла.")


def external_id(row: dict[str, Any]) -> str:
    source = "|".join(
        str(row[key])
        for key in ("created_at", "device_number", "card_number", "address", "volume", "cost")
    )
    return "cr-excel-" + hashlib.sha256(source.encode()).hexdigest()[:32]


def transaction(row: dict[str, Any], point: tuple[float, float]) -> dict[str, Any]:
    description_parts = [
        f"Поставщик: {row['supplier']}" if row["supplier"] else None,
        f"АЗС: {row['station']}" if row["station"] else None,
    ]
    return {
        "deviceNumber": row["device_number"],
        "cardNumber": row["card_number"],
        "externalId": external_id(row),
        "createdAt": row["created_at"],
        "volume": row["volume"],
        "cost": row["cost"],
        "description": "; ".join(part for part in description_parts if part),
        "address": row["address"],
        "x": point[0],
        "y": point[1],
    }


def build_transactions(api: httpx.Client, rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    cache: dict[str, tuple[float, float]] = {}
    transactions: list[dict[str, Any]] = []
    for row in rows:
        address = row["address"]
        if address not in cache:
            cache[address] = coordinates(api, address)
        transactions.append(transaction(row, cache[address]))
    return transactions


def main() -> int:
    parser = argparse.ArgumentParser(description="Импорт транзакций заправки в CargoRun.")
    parser.add_argument("--file", required=True, type=Path, help="Путь к Excel-файлу.")
    parser.add_argument("--device-id", required=True, help="X-Device-Id из браузерного запроса CargoRun.")
    parser.add_argument("--limit", type=int, default=10, help="Число первых строк для обработки (по умолчанию 10).")
    parser.add_argument("--apply", action="store_true", help="Отправить транзакции. Без флага выполняется проверка.")
    args = parser.parse_args()
    if args.limit < 1:
        parser.error("--limit должен быть положительным числом.")
    if not args.file.is_file():
        parser.error(f"Файл не найден: {args.file}")

    token = getpass.getpass("Bearer-токен CargoRun: ").removeprefix("Bearer ").strip()
    if not token:
        parser.error("Токен не указан.")

    rows = report_rows(args.file, args.limit)
    with client(token, args.device_id) as api:
        transactions = build_transactions(api, rows)
        print(f"Подготовлено транзакций: {len(transactions)}.")
        if not args.apply:
            print("Проверка завершена. Для загрузки добавьте --apply.")
            return 0
        response = api.post("/api/RefuelingTransactions/Create", json={"transactions": transactions})
        response.raise_for_status()
        print(f"CargoRun принял пакет из {len(transactions)} транзакций (HTTP {response.status_code}).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
