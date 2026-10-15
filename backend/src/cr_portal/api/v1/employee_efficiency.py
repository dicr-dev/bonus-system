from datetime import date
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font

from cr_portal.api.deps import admin_user, db_session
from cr_portal.services.app_settings import get_business_settings
from cr_portal.services.employee_efficiency import employee_efficiency_report

router = APIRouter()


async def _report(session, date_from, date_to, employee_id, module, integration_1c):
    if employee_id and employee_id != "unassigned":
        try:
            UUID(employee_id)
        except ValueError as error:
            raise HTTPException(422, "Некорректный сотрудник") from error
    business = await get_business_settings(session)
    try:
        return await employee_efficiency_report(session, business, date_from=date_from, date_to=date_to,
                                               employee_id=employee_id, module=module, integration_1c=integration_1c)
    except ValueError as error:
        raise HTTPException(422, str(error)) from error


@router.get("/employee-efficiency")
async def report(date_from: date | None = Query(None), date_to: date | None = Query(None),
                 employee_id: str | None = Query(None), module: str | None = Query(None),
                 integration_1c: bool | None = Query(None), session=Depends(db_session), _admin=Depends(admin_user)):
    return await _report(session, date_from, date_to, employee_id, module, integration_1c)


def _cell(value):
    # Spreadsheet applications must treat CRM titles and stage names as text.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def efficiency_workbook(data):
    book = Workbook()
    info = book.active
    info.title = "Период и правила"
    info.append(["Отчёт", "Эффективность сотрудников"])
    info.append(["Период с", data["date_from"] or "За весь срок"])
    info.append(["Период по", data["date_to"]])
    info.append(["Ответственность", "По текущему ответственному за внедрение"])
    info.append(["Длительность", "Полный срок сделки от создания до завершения; для незавершённых — до конца периода"])
    info.append(["Доля успешных", "Успешные / (успешные + неуспешные закрытия в периоде)"])
    for warning in data["warnings"]:
        info.append(["Замечание", warning])
    metrics = [("new_count", "Новые"), ("work_count", "Были в работе"), ("active_end_count", "Открыты на конец периода"),
               ("won_count", "Успешные"), ("lost_count", "Неуспешные"), ("won_amount", "Сумма успешных"),
               ("lost_amount", "Сумма неуспешных"), ("success_rate", "Доля успешных, %"),
               ("avg_completion_days", "Средний полный срок успешных, дни"), ("with_1c_count", "С интеграцией 1С"),
               ("avg_days_with_1c", "Средний срок с 1С"), ("avg_days_without_1c", "Средний срок без 1С"),
               ("first_training_count", "Первое обучение"), ("second_training_count", "Второе обучение"),
               ("reports_training_count", "Обучение отчётам"), ("training_count", "Всего обучений"),
               ("fully_trained_count", "Все обучения завершены"), ("support_created_count", "Создано сделок Сопровождения")]
    staff = book.create_sheet("Сотрудники")
    modules = book.create_sheet("Модули")
    staff.append(["Сотрудник", "Отдел", "Воронка", *[label for _, label in metrics]])
    modules.append(["Сотрудник", "Модуль", "Воронка", *[label for _, label in metrics]])
    funnel_names = {"tech_integration": "Техинтеграция", "implementation": "Внедрение"}
    for person in data["employees"]:
        for funnel, name in funnel_names.items():
            staff.append([_cell(person["name"]), _cell(person["department"]), name,
                          *[float(person["metrics"][funnel][key]) if key in {"won_amount", "lost_amount"} else person["metrics"][funnel][key] for key, _ in metrics]])
        for module in person["modules"]:
            for funnel, name in funnel_names.items():
                modules.append([_cell(person["name"]), _cell(module["module_name"]), name,
                                *[float(module["metrics"][funnel][key]) if key in {"won_amount", "lost_amount"} else module["metrics"][funnel][key] for key, _ in metrics]])
    detail = book.create_sheet("Сделки")
    detail.append(["Сотрудник", "ID сделки", "Название", "Воронка", "Модуль", "Интеграция 1С", "Сумма", "Создана", "Закрыта",
                   "Статус на конец периода", "Текущая стадия", "Полный срок, дни", "Причина отказа", "Первое обучение", "Второе обучение",
                   "Обучение отчётам", "Обучений пройдено", "Требуется обучений", "История загружена"])
    for person in data["employees"]:
        for row in person["deals"]:
            detail.append([_cell(person["name"]), row["bitrix_id"], _cell(row["title"]), funnel_names[row["funnel"]], _cell(row["module_name"]),
                           "Да" if row["integration_1c"] else "Нет", float(row["amount"]), row["created_at"], row["closed_at"], row["status_at_end"],
                           _cell(row["current_stage_title"]), row["duration_days"], _cell(row["loss_reason"]),
                           *[row["training_dates"].get(name) for name in ("first", "second", "reports")],
                           row["training_done"], row["training_target"], "Да" if row["history_available"] else "Нет"])
    losses = book.create_sheet("Причины отказов")
    losses.append(["Воронка", "Причина", "Количество", "Сумма"])
    for row in data["losses"]:
        losses.append([funnel_names[row["funnel"]], _cell(row["reason"]), row["count"], float(row["amount"])])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    return book


@router.get("/employee-efficiency/export")
async def export(date_from: date | None = Query(None), date_to: date | None = Query(None),
                 employee_id: str | None = Query(None), module: str | None = Query(None),
                 integration_1c: bool | None = Query(None), session=Depends(db_session), _admin=Depends(admin_user)):
    data = await _report(session, date_from, date_to, employee_id, module, integration_1c)
    buffer = BytesIO()
    efficiency_workbook(data).save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="employee-efficiency.xlsx"'})
