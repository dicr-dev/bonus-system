"""One-time fuel-card binding import.  Tokens are intentionally never persisted."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

import httpx
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

FUEL_CARDS_APPLY_URL = "https://delko.cargorun.ru/api/fuelCards/apply"
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 5_000

BRANDS = {
    "таиф": "Tatneft", "tatneft": "Tatneft",
    "делкойл": "Azimuth", "azimuth": "Azimuth",
    "газпромнефть": "Gazpromneft", "gazpromneft": "Gazpromneft",
    "роснефть": "Rosneft", "rosneft": "Rosneft",
}
PLATE_COLUMNS = (
    "госномер", "гос номер", "государственный номер", "номер", "номер тс", "номер автомобиля",
    "номер машины", "регистрационный номер", "рег номер", "vehicle plate", "plate", "reg number",
)
ID_COLUMNS = (
    "id", "car id", "id тс", "ид тс", "идентификатор тс", "vehicle id", "id автомобиля",
    "id машины", "ид автомобиля", "ид машины", "идентификатор автомобиля", "идентификатор машины",
)


@dataclass
class Binding:
    row: int
    number: str
    brand: str
    plate: str
    car_id: int | None
    error: str | None = None


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value).casefold()).strip()


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-zа-я0-9]", "", _key(value))


def _plate(value: Any) -> str:
    return re.sub(r"[^0-9A-ZА-Я]", "", _text(value).upper())


def _brand_column(header: str) -> str | None:
    normalized = _header_key(header)
    for label, brand in BRANDS.items():
        if _header_key(label) in normalized:
            return brand
    return None


def _contains_column(headers: list[str], names: tuple[str, ...]) -> bool:
    candidates = {_header_key(name) for name in names}
    return bool(candidates & {_header_key(header) for header in headers})


def _is_import_header(headers: list[str]) -> bool:
    return (
        _contains_column(headers, PLATE_COLUMNS)
        and _contains_column(headers, ID_COLUMNS)
        and any(_brand_column(header) for header in headers)
    )


def _rows(content: bytes, label: str) -> tuple[list[tuple[int, dict[str, Any]]], str | None]:
    if not content or len(content) > MAX_FILE_SIZE:
        return [], f"Файл «{label}» пуст или превышает 10 МБ."
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = book.active
        values = sheet.iter_rows(values_only=True)
        header_values = None
        for candidate_row_number, row in enumerate(values, start=1):
            candidate = [_key(value) for value in row]
            if _is_import_header(candidate):
                header_values = candidate
                break
            if candidate_row_number >= 20:
                break
        if header_values is None:
            return [], f"В файле «{label}» нет заголовков."
        keys = header_values
        result: list[tuple[int, dict[str, Any]]] = []
        for row_number, row in enumerate(values, start=candidate_row_number + 1):
            if len(result) >= MAX_ROWS:
                return [], f"В файле «{label}» больше {MAX_ROWS} строк."
            result.append((row_number, {keys[index]: value for index, value in enumerate(row) if index < len(keys)}))
        return result, None
    except (BadZipFile, InvalidFileException, OSError, TypeError, ValueError):
        return [], f"Не удалось прочитать «{label}». Нужен файл Excel .xlsx без пароля."


def _value(row: dict[str, Any], names: tuple[str, ...]) -> Any:
    keys = {_header_key(name) for name in names}
    for header, value in row.items():
        if _header_key(header) in keys and _text(value):
            return value
    return None


def preview(import_file: bytes) -> dict[str, Any]:
    rows, error = _rows(import_file, "импорт")
    if error:
        raise ValueError(error)
    if not rows:
        raise ValueError("В файле нет строк.")
    headers = rows[0][1].keys()
    card_columns = {header: brand for header in headers if (brand := _brand_column(header))}
    if not card_columns:
        raise ValueError("Не найдены колонки карт: ТАИФ, Делкойл, Газпромнефть или Роснефть.")
    if not _contains_column(list(headers), PLATE_COLUMNS) or not _contains_column(list(headers), ID_COLUMNS):
        raise ValueError("В файле нужны колонки «госномер» и «id» ТС.")
    bindings: list[Binding] = []
    seen_cards: set[str] = set()
    for row_number, row in rows:
        plate = _plate(_value(row, PLATE_COLUMNS))
        try:
            car_id = int(_text(_value(row, ID_COLUMNS)))
        except ValueError:
            car_id = None
        for header, brand in card_columns.items():
            number = _text(row.get(header))
            if not number:
                continue
            error = None
            if not plate or car_id is None:
                error = "Не указан госномер или id ТС."
            elif number in seen_cards:
                error = "Повторяющийся номер карты."
            seen_cards.add(number)
            bindings.append(Binding(row=row_number, number=number, brand=brand, plate=plate, car_id=car_id, error=error))
    if not bindings:
        raise ValueError("В брендовых колонках нет номеров топливных карт.")
    valid = [item for item in bindings if item.error is None]
    return {"total": len(bindings), "ready": len(valid), "errors": len(bindings) - len(valid), "items": [asdict(item) for item in bindings]}


def binding_key(item: dict[str, Any]) -> str:
    return f"{item['row']}:{item['brand']}:{item['number']}"


def selected_bindings(bindings: list[dict[str, Any]], selected_keys: set[str]) -> list[dict[str, Any]]:
    selected = [item for item in bindings if binding_key(item) in selected_keys]
    if not selected:
        raise ValueError("Выберите хотя бы одну корректную карту.")
    return selected


async def apply(bindings: list[dict[str, Any]], token: str) -> dict[str, Any]:
    if not token.strip():
        raise ValueError("Укажите временный токен CargoRun.")
    ready = [item for item in bindings if not item.get("error") and item.get("car_id")]
    results: list[dict[str, Any]] = []
    headers = {"Authorization": f"Bearer {token.strip()}", "Accept": "application/json"}
    async with httpx.AsyncClient(timeout=30) as client:
        for item in ready:
            payload = {"isValid": True, "number": item["number"], "type": item["brand"], "hasNoLimitManagement": True, "carId": item["car_id"]}
            try:
                response = await client.post(FUEL_CARDS_APPLY_URL, headers=headers, json=payload)
                if response.is_success:
                    results.append({"row": item["row"], "number": item["number"], "brand": item["brand"], "status": "applied"})
                else:
                    results.append({"row": item["row"], "number": item["number"], "brand": item["brand"], "status": "error", "message": f"CargoRun вернул HTTP {response.status_code}."})
            except httpx.HTTPError:
                results.append({"row": item["row"], "number": item["number"], "brand": item["brand"], "status": "error", "message": "Не удалось связаться с CargoRun."})
    return {"applied": sum(item["status"] == "applied" for item in results), "failed": sum(item["status"] == "error" for item in results), "items": results}
