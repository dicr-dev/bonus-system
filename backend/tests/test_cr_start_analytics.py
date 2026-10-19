import json
import unittest
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace as Obj
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook

from cr_portal.api.deps import current_user, db_session
from cr_portal.api.v1.cr_start_analytics import router, workbook
from cr_portal.services.cr_start_analytics import build_cr_start_analytics, module_name


def dt(value):
    return datetime.fromisoformat(value + "+03:00")


def user(name="Manager", bitrix=1, active=True):
    return Obj(id=uuid4(), bitrix_id=bitrix, full_name=name, is_active=active, department_name="Отдел внедрения")


def deal(owner=None, **values):
    return Obj(**(dict(id=uuid4(), bitrix_id=100, funnel="cr_start", category_id=4,
                      implementation_responsible_user_id=owner.id if owner else None,
                      title="Company", stage_title="Период оплачен", stage_id="paid", module_name="Логистика",
                      opportunity=1000, company_bitrix_id=42, company_name="Company",
                      created_time=dt("2026-01-01T00:00:00"), raw_json=json.dumps({"START": "2026-01-02"})) | values))


def event(d, date, stage, title, id=1):
    return Obj(deal_id=d.id, occurred_at=dt(date + "T00:00:00"), stage_id=stage,
               stage_title=title, bitrix_event_id=id, category_id=d.category_id)


def history(d):
    return [event(d,"2026-01-01","new","Новая",1),
            event(d,"2026-01-02","commercial","Коммерческое использование",2),
            event(d,"2026-02-02","lost","Слишком дорого",3),
            event(d,"2026-03-02","commercial","Коммерческое использование",4),
            event(d,"2026-04-02","paid","Период оплачен",5)]


def task(id=10, deals=(100,), **values):
    return Obj(**(dict(bitrix_id=id,title="Task",crm_deal_ids_json=json.dumps(deals),deadline=None,
                      status=2,closed_time=None) | values))


def elapsed(id=1, task_id=10, author=1, when="2026-01-01T12:00:00", seconds=3600):
    return Obj(bitrix_id=id,task_bitrix_id=task_id,user_bitrix_id=author,created_time=dt(when),seconds=seconds)


def build(deals, users=(), events=(), tasks=(), logs=(), **kw):
    return build_cr_start_analytics(deals, users, events, tasks, logs, "START", now=dt("2026-05-02T00:00:00"), **kw)


class AnalyticsTests(unittest.TestCase):
    def test_reconnections_exclude_break_and_commercial_stage_changes(self):
        u=user()
        d=deal(u)
        data=build([d],[u],history(d))
        row=data["deals"][0]
        self.assertTrue(row["history_reliable"])
        self.assertEqual(row["usage_days"],92)
        self.assertEqual(row["reconnections"],1)
        self.assertEqual(len(row["periods"]),2)
        self.assertEqual(row["retention"],{"1":False,"3":True,"6":None})
        self.assertEqual(data["totals"]["retention_6_sample"],0)
        self.assertEqual(data["totals"]["commercial_count"],1)
        self.assertEqual(row["stage_days"],30)

    def test_lost_after_is_not_erased_from_historical_conversion(self):
        u=user()
        d=deal(u,stage_id="lost",stage_title="Слишком дорого")
        events=history(d)[:3]
        data=build([d],[u],events)
        self.assertEqual(data["totals"]["ever_commercial_count"],1)
        self.assertEqual(data["totals"]["commercial_count"],0)
        self.assertEqual(data["totals"]["lost_after_count"],1)
        self.assertEqual(data["totals"]["conversion"],100)
        self.assertEqual(data["deals"][0]["usage_days"],31)

    def test_lost_before_date_comes_from_entry_into_loss_group_not_reason_edits(self):
        d=deal(stage_id="lost2",stage_title="Перестали отвечать",raw_json="{}")
        events=[event(d,"2026-01-01","new","Новая",1),event(d,"2026-01-04","lost","Слишком дорого",2),
                event(d,"2026-02-04","lost2","Перестали отвечать",3)]
        data=build([d],events=events)
        self.assertTrue(data["deals"][0]["loss_at"].startswith("2026-01-04"))
        self.assertEqual(data["monthly"][0]["lost_before"],1)
        self.assertEqual(sum(m["lost_before"] for m in data["monthly"]),1)

    def test_no_history_never_fabricates_churn_or_usage(self):
        d=deal(stage_id="lost",stage_title="Слишком дорого")
        data=build([d])
        row=data["deals"][0]
        self.assertIsNone(row["usage_days"])
        self.assertIsNone(row["loss_at"])
        self.assertIsNone(row["stage_days"])
        self.assertFalse(row["history_reliable"])
        self.assertIsNone(data["totals"]["retention_1"])

    def test_missing_field_is_not_silently_replaced_by_history(self):
        d=deal(raw_json="{}")
        row=build([d],events=history(d))["deals"][0]
        self.assertTrue(row["reached"])
        self.assertIsNone(row["first_commercial_at"])
        self.assertIsNone(row["usage_days"])

    def test_snapshot_mismatch_prevents_reliable_usage(self):
        d=deal()
        row=build([d],events=history(d)[:3])["deals"][0]
        self.assertFalse(row["history_reliable"])
        self.assertIsNone(row["usage_days"])

    def test_same_day_loss_after_launch_is_a_real_period_boundary(self):
        d=deal(stage_id="lost",stage_title="Слишком дорого")
        events=history(d)[:2]
        lost=event(d,"2026-01-02","lost","Слишком дорого",3)
        lost.occurred_at=dt("2026-01-02T16:00:00")
        events.append(lost)
        row=build([d],events=events)["deals"][0]
        self.assertTrue(row["history_reliable"])
        self.assertEqual(row["usage_days"],0.67)

    def test_ambiguous_same_day_loss_before_first_observed_launch_is_flagged(self):
        d=deal(stage_id="commercial",stage_title="Коммерческое использование")
        lost=event(d,"2026-01-02","lost","Слишком дорого",2)
        lost.occurred_at=dt("2026-01-02T10:00:00")
        launch=event(d,"2026-01-02","commercial","Коммерческое использование",3)
        launch.occurred_at=dt("2026-01-02T13:00:00")
        row=build([d],events=[event(d,"2026-01-01","new","Новая"),lost,launch])["deals"][0]
        self.assertFalse(row["history_reliable"])
        self.assertIsNone(row["usage_days"])

    def test_late_history_and_future_start_have_no_usage_estimate(self):
        d=deal()
        self.assertIsNone(build([d],events=history(d)[2:])["deals"][0]["usage_days"])
        d.raw_json=json.dumps({"START":"2027-01-01"})
        self.assertIsNone(build([d],events=history(d))["deals"][0]["first_commercial_at"])

    def test_duplicates_are_not_excluded_from_conversion_and_companies_are_unique(self):
        first=deal()
        duplicate=deal(bitrix_id=101,stage_title="Дубль",stage_id="lost",raw_json="{}")
        data=build([first,duplicate],events=history(first))
        self.assertEqual(data["totals"]["conversion"],50)
        self.assertEqual(data["totals"]["companies_count"],1)
        self.assertEqual(data["totals"]["lost_before_count"],1)

    def test_module_mapping_and_unrecognized_stages(self):
        self.assertEqual(module_name(deal(module_name="КР Старт ТМ")),"Топливный модуль")
        self.assertEqual(module_name(deal(module_name="КР Старт Эксплуатация")),"Эксплуатация")
        self.assertEqual(module_name(deal(module_name="Иное")),"Логистика")
        data=build([deal(stage_title="Не учитывается")])
        self.assertEqual(data["totals"]["deals_count"],0)
        self.assertEqual(data["excluded_count"],1)

    def test_hours_remain_with_author_not_current_owner_and_split_by_phase(self):
        former,current=user("Former",1,False),user("Current",2)
        d=deal(current)
        logs=[elapsed(1,author=1),elapsed(2,author=1,when="2026-01-10T12:00:00"),
              elapsed(3,author=1,when="2026-02-10T12:00:00"),elapsed(4,author=2)]
        data=build([d],[former,current],history(d),[task()],logs)
        people={m["name"]:m for m in data["managers"]}
        self.assertEqual(people["Former"]["author_hours"],3)
        self.assertEqual(people["Former"]["deals_count"],0)
        self.assertEqual(people["Former"]["before_hours"],1)
        self.assertEqual(people["Former"]["commercial_hours"],1)
        self.assertEqual(people["Former"]["break_hours"],1)
        self.assertEqual(people["Current"]["author_hours"],1)
        self.assertEqual(people["Current"]["hours"],4)
        self.assertEqual(data["deals"][0]["hours"],4)

    def test_multiple_deal_task_is_excluded_no_double_counting(self):
        u=user()
        d=deal(u)
        data=build([d],[u],history(d),[task(deals=[100,101])],[elapsed()])
        self.assertEqual(len(data["issues"]),1)
        self.assertIsNone(data["deals"][0]["hours"])
        self.assertEqual(data["deals"][0]["tasks"],[])

    def test_no_logs_null_hours_zero_staff_and_missing_owner(self):
        u=user()
        d=deal(stage_title="Новая",stage_id="new",raw_json="{}")
        data=build([d],[u],tasks=[task()])
        self.assertIsNone(data["deals"][0]["hours"])
        manager=next(m for m in data["managers"] if m["id"]==str(u.id))
        self.assertEqual(manager["deals_count"],0)
        self.assertIsNone(manager["author_hours"])
        self.assertNotIn("unassigned",[m["id"] for m in data["managers"]])
        self.assertEqual(data["deals"][0]["owner_id"],"unassigned")

    def test_only_implementation_staff_in_list_active_before_inactive(self):
        active=user("Z active",1)
        inactive=user("A inactive",2,False)
        other=user("Other department",3)
        other.department_name="Разработка 1С"
        unknown=user("No department",4)
        unknown.department_name=None
        d=deal(other)
        data=build([d],[inactive,other,unknown,active],history(d),[task()],[elapsed(author=3)])
        self.assertEqual([m["id"] for m in data["managers"]],[str(active.id),str(inactive.id)])
        self.assertEqual(data["deals"][0]["hours"],1)
        self.assertEqual(data["deals"][0]["effort"][0]["author"],other.full_name)

    def test_multiple_departments_and_zero_inactive_staff_are_included(self):
        a=user("B active",1)
        b=user("A active",2)
        b.department_name="Компания; Отдел внедрения; Другой отдел"
        retired=user("Inactive without deals",3,False)
        data=build([], [a,retired,b])
        self.assertEqual([m["name"] for m in data["managers"]],["A active","B active","Inactive without deals"])
        self.assertTrue(all(m["deals_count"]==0 for m in data["managers"]))

    def test_duplicate_logs_and_future_entries_are_not_counted(self):
        d=deal()
        e=elapsed()
        data=build([d],events=history(d),tasks=[task()],logs=[e,e,elapsed(2,when="2027-01-01T00:00:00")])
        self.assertEqual(data["deals"][0]["hours"],1)

    def test_seconds_are_summed_before_rounding_and_lost_effort_is_separate(self):
        u=user()
        deals=[deal(u,bitrix_id=100+i,stage_title="Дубль",stage_id="lost",raw_json="{}") for i in range(36)]
        tasks=[task(id=10+i,deals=[100+i]) for i in range(36)]
        logs=[elapsed(id=1+i,task_id=10+i,seconds=1) for i in range(36)]
        data=build(deals,[u],tasks=tasks,logs=logs)
        self.assertEqual(data["totals"]["hours"],0.01)
        self.assertEqual(data["totals"]["lost_before_hours"],0.01)
        self.assertEqual(data["managers"][0]["lost_before_author_hours"],0.01)

    def test_only_open_overdue_tasks_count(self):
        d=deal()
        deadline=dt("2026-01-01T00:00:00")
        data=build([d],events=history(d),tasks=[task(deadline=deadline),task(id=11,deadline=deadline,status=5)])
        self.assertEqual(data["deals"][0]["overdue_tasks"],1)

    def test_filters_scope_deals_not_ownership_of_time(self):
        u,v=user("U",1),user("V",2)
        d=deal(u,module_name="КР Старт ТМ")
        other=deal(v,bitrix_id=101)
        data=build([d,other],[u,v],history(d),[task()],[elapsed(author=2)],employee_id=str(u.id),module="Топливный модуль")
        self.assertEqual(len(data["deals"]),1)
        self.assertEqual([m["id"] for m in data["managers"]],[str(u.id)])
        self.assertEqual(data["deals"][0]["effort"][0]["author_id"],str(v.id))
        self.assertEqual(data["deals"][0]["effort"][0]["hours"],1)

    def test_cohorts_and_dynamics_follow_event_months(self):
        d=deal()
        data=build([d],events=history(d))
        months={m["month"]:m for m in data["monthly"]}
        self.assertEqual(months["2026-01"]["first_launch"],1)
        self.assertEqual(months["2026-02"]["departures"],1)
        self.assertEqual(months["2026-03"]["reconnections"],1)
        self.assertEqual(data["cohorts"][0]["deals_count"],1)

    def test_workbook_preserves_numbers_and_escapes_crm_formulas(self):
        d=deal(title="=HYPERLINK(1)")
        data=build([d],events=history(d))
        book=workbook(data)
        buffer=BytesIO()
        book.save(buffer)
        buffer.seek(0)
        restored=load_workbook(buffer)
        self.assertEqual(restored["Сделки"].cell(2,2).value,"'=HYPERLINK(1)")
        self.assertEqual(restored["Сделки"].cell(2,7).data_type,"n")
        self.assertIn("Списания времени",restored.sheetnames)


class AccessTests(unittest.IsolatedAsyncioTestCase):
    def app(self, admin):
        app=FastAPI()
        app.include_router(router,prefix="/analytics")
        app.dependency_overrides[current_user]=lambda: Obj(is_admin=admin)
        app.dependency_overrides[db_session]=lambda: None
        return app

    async def test_employee_cannot_read_or_export(self):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app(False)),base_url="http://test") as client:
            for path in ("/analytics/cr-start","/analytics/cr-start/export"):
                self.assertEqual((await client.get(path)).status_code,403)

    async def test_admin_and_invalid_module(self):
        with patch("cr_portal.api.v1.cr_start_analytics.get_business_settings",AsyncMock(return_value=Obj())), \
             patch("cr_portal.api.v1.cr_start_analytics.cr_start_analytics",AsyncMock(return_value=build([]))) as report:
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app(True)),base_url="http://test") as client:
                self.assertEqual((await client.get("/analytics/cr-start")).status_code,200)
                self.assertEqual((await client.get("/analytics/cr-start",params={"module":"bad"})).status_code,422)
                report.assert_awaited_once()
