import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook

from cr_portal.api.deps import current_user, db_session
from cr_portal.api.v1.analytics import router


class WeeklyOvAccessTests(unittest.IsolatedAsyncioTestCase):
    def app(self, department=None, admin=False, authenticated=True):
        app = FastAPI()
        app.include_router(router, prefix="/analytics")
        async def user():
            if not authenticated:
                raise HTTPException(401, "Authentication required")
            return SimpleNamespace(is_admin=admin, department_name=department)
        app.dependency_overrides[current_user] = user
        app.dependency_overrides[db_session] = lambda: None
        return app

    async def request_report(self, app, expected):
        with patch("cr_portal.api.v1.analytics.weekly_ov_report", AsyncMock(return_value=[])) as report:
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                for path in ("/weekly-ov", "/weekly-ov/export"):
                    response = await client.get("/analytics"+path,params={"date_from":"2026-10-19","date_to":"2026-10-25"})
                    self.assertEqual(response.status_code, expected, response.text[:100])
                    if expected == 200 and path.endswith("export"):
                        workbook = load_workbook(BytesIO(response.content))
                        self.assertEqual(workbook.active.cell(1,1).value, "Воронка")
            self.assertEqual(report.await_count, 2 if expected == 200 else 0)

    async def test_implementation_employee_can_view_and_export(self):
        await self.request_report(self.app("Отдел внедрения"),200)

    async def test_multiple_departments_can_view_and_export(self):
        await self.request_report(self.app("Отдел сопровождения; Отдел внедрения"),200)

    async def test_admin_can_view_and_export_without_department(self):
        await self.request_report(self.app(admin=True),200)

    async def test_other_departments_and_partial_matches_are_denied(self):
        for department in (None,"","Разработка 1С","Отдел сопровождения","Отдел внедрения (другой)"):
            with self.subTest(department=department):
                await self.request_report(self.app(department),403)

    async def test_unauthenticated_is_denied(self):
        await self.request_report(self.app(authenticated=False),401)

    async def test_other_analytics_remain_administrator_only(self):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app("Отдел внедрения")),base_url="http://test") as client:
            for path in ("/deal-groups","/deals-in-work"):
                self.assertEqual((await client.get("/analytics"+path)).status_code,403)

    async def test_invalid_dates_still_rejected(self):
        with patch("cr_portal.api.v1.analytics.weekly_ov_report", AsyncMock()) as report:
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app("Отдел внедрения")),base_url="http://test") as client:
                for path in ("/weekly-ov","/weekly-ov/export"):
                    response=await client.get("/analytics"+path,params={"date_from":"2026-10-25","date_to":"2026-10-19"})
                    self.assertEqual(response.status_code,422)
                report.assert_not_awaited()


if __name__ == "__main__":
    unittest.main()
