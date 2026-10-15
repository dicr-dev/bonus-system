import json
import unittest
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook

from cr_portal.api.v1.employee_efficiency import efficiency_workbook
from cr_portal.services.employee_efficiency import build_employee_efficiency


class EmployeeEfficiencyTests(unittest.TestCase):
    def setUp(self):
        self.owner = SimpleNamespace(id=uuid4(), full_name="Сотрудник", department_name="Отдел внедрения", is_active=True)
        self.business = SimpleNamespace(field_source_deal_id="source", field_first_training_date="first", field_second_training_date="second", field_reports_training_date="reports")
        self.now = datetime(2026, 10, 14, 12, tzinfo=UTC)

    def deal(self, bid=1, **kwargs):
        values = dict(id=uuid4(), bitrix_id=bid, title=f"Сделка {bid}", funnel="implementation", category_id=7,
                      implementation_responsible_user_id=self.owner.id, module_name="Логистика", integration_1c=False,
                      created_time=datetime(2026, 8, 1, tzinfo=UTC), closed_time=None, status="in_progress",
                      stage_title="В работе", opportunity=Decimal("1000"), raw_json="{}")
        values.update(kwargs)
        return SimpleNamespace(**values)

    def event(self, deal, when, semantic, title="Стадия", event_id=1):
        return SimpleNamespace(deal_id=deal.id, category_id=deal.category_id, occurred_at=when,
                               semantic=semantic, stage_title=title, stage_id="STAGE", bitrix_event_id=event_id)

    def report(self, deals, history=(), **kwargs):
        return build_employee_efficiency(deals, [self.owner], history, self.business,
                                         date_from=date(2026, 9, 1), date_to=date(2026, 9, 30), now=self.now, **kwargs)

    def test_period_counts_carryover_closures_and_full_lifetime_duration(self):
        won = self.deal(1, status="won", closed_time=datetime(2026, 9, 10, tzinfo=UTC))
        lost = self.deal(2, status="lost", stage_title="Слишком дорого", closed_time=datetime(2026, 9, 15, tzinfo=UTC))
        new = self.deal(3, created_time=datetime(2026, 9, 2, tzinfo=UTC))
        old = self.deal(4, status="won", closed_time=datetime(2026, 8, 20, tzinfo=UTC))
        report = self.report([won, lost, new, old])
        stats = report["summary"]["implementation"]
        self.assertEqual((stats["new_count"], stats["work_count"], stats["won_count"], stats["lost_count"], stats["active_end_count"]), (1, 3, 1, 1, 1))
        self.assertEqual(stats["avg_completion_days"], 40)
        self.assertEqual(stats["success_rate"], 50)
        self.assertEqual(report["losses"][0]["reason"], "Слишком дорого")
        self.assertEqual(report["total_deals"], 3)

    def test_future_closure_and_moscow_midnight_boundary(self):
        future = self.deal(1, status="lost", closed_time=datetime(2026, 10, 5, tzinfo=UTC))
        after_period = self.deal(2, created_time=datetime(2026, 9, 30, 21, tzinfo=UTC))
        stats = self.report([future, after_period])["summary"]["implementation"]
        self.assertEqual(stats["active_end_count"], 1)
        self.assertEqual(stats["lost_count"], 0)
        self.assertEqual(stats["deal_count"], 1)

    def test_reopen_history_is_used_instead_of_current_status(self):
        deal = self.deal(status="lost", closed_time=datetime(2026, 10, 5, tzinfo=UTC))
        events = [self.event(deal, datetime(2026, 8, 20, tzinfo=UTC), "S", event_id=1),
                  self.event(deal, datetime(2026, 9, 5, tzinfo=UTC), "P", event_id=2),
                  self.event(deal, datetime(2026, 10, 5, tzinfo=UTC), "F", event_id=3)]
        report = self.report([deal], events)
        stats = report["summary"]["implementation"]
        self.assertEqual((stats["active_end_count"], stats["won_count"], stats["lost_count"]), (1, 0, 0))
        self.assertEqual(report["history_deals"], 1)
        self.assertFalse(report["warnings"])

    def test_terminal_stage_edits_do_not_move_closure_into_another_month(self):
        deal = self.deal(status="lost", closed_time=None)
        events = [self.event(deal, datetime(2026, 8, 20, tzinfo=UTC), "F", event_id=1),
                  self.event(deal, datetime(2026, 9, 10, tzinfo=UTC), "F", title="Уточнённая причина", event_id=2)]
        self.assertEqual(self.report([deal], events)["total_deals"], 0)

    def test_trainings_are_module_specific_and_filtered_by_period(self):
        dates = json.dumps({"first": "2026-08-10", "second": "2026-09-05", "reports": "2026-09-10"})
        logistics = self.deal(1, raw_json=dates)
        fuel = self.deal(2, module_name="Топливный модуль", raw_json=dates)
        stats = self.report([logistics, fuel])["summary"]["implementation"]
        self.assertEqual(stats["training_count"], 2)
        self.assertEqual(stats["first_training_count"], 0)
        self.assertEqual(stats["fully_trained_count"], 2)
        details = self.report([fuel])["employees"][0]["deals"][0]
        self.assertEqual((details["training_target"], details["training_done"]), (1, 1))

    def test_support_transfer_and_1c_are_linked_to_implementation_owner(self):
        tech = self.deal(1, funnel="tech_integration", category_id=1, integration_1c=True)
        implementation = self.deal(2, raw_json=json.dumps({"source": ["D_1"]}))
        support = self.deal(3, funnel="support", category_id=18, implementation_responsible_user_id=uuid4(),
                            created_time=datetime(2026, 9, 5, tzinfo=UTC), raw_json=json.dumps({"source": ["D_1", "D_2"]}))
        report = self.report([tech, implementation, support], integration_1c=True)
        self.assertEqual(report["summary"]["implementation"]["support_created_count"], 1)
        self.assertEqual(report["summary"]["implementation"]["with_1c_count"], 1)
        self.assertEqual(len(report["employees"]), 1)

    def test_inactive_current_owner_and_unassigned_deals_are_not_lost(self):
        self.owner.is_active = False
        report = self.report([self.deal(1), self.deal(2, implementation_responsible_user_id=None)])
        self.assertEqual(len(report["employees"]), 2)
        self.assertTrue(any(person["id"] == "unassigned" for person in report["employees"]))
        self.assertEqual(report["summary"]["implementation"]["work_count"], 2)

    def test_missing_close_date_is_not_counted_as_a_period_completion(self):
        deal = self.deal(status="won", closed_time=None, created_time=datetime(2026, 9, 3, tzinfo=UTC))
        report = self.report([deal])
        self.assertEqual(report["summary"]["implementation"]["won_count"], 0)
        self.assertTrue(any("нет даты закрытия" in text for text in report["warnings"]))

    def test_all_time_and_invalid_range(self):
        report = build_employee_efficiency([self.deal(status="won", closed_time=datetime(2026, 8, 20, tzinfo=UTC))],
                                          [self.owner], [], self.business, now=self.now)
        self.assertEqual(report["summary"]["implementation"]["won_count"], 1)
        with self.assertRaises(ValueError):
            build_employee_efficiency([], [], [], self.business,
                date_from=date(2026, 10, 5), date_to=date(2026, 9, 1), now=self.now)

    def test_excel_preserves_numbers_and_does_not_execute_crm_titles(self):
        report = self.report([self.deal(title="=1+1")])
        workbook = efficiency_workbook(report)
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        saved = load_workbook(buffer)
        self.assertEqual(saved["Сделки"]["C2"].value, "'=1+1")
        self.assertEqual(saved["Сделки"]["G2"].value, 1000)
        self.assertEqual(len(saved.sheetnames), 5)


if __name__ == "__main__":
    unittest.main()
